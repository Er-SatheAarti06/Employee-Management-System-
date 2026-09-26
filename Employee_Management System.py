import tkinter as tk
from tkinter import messagebox
import openpyxl
import os


# ---------------- LOGIN ----------------

def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "Aarti" and password == "Sathe @1234":
        messagebox.showinfo("Login", "Login Successful")
        login_window.destroy()
        open_dashboard()
    else:
        messagebox.showerror("Login", "Invalid Username or Password")


# ---------------- ADD RECORD ----------------

def add_record():
    add_window = tk.Toplevel()
    add_window.title("Add Employee Record")
    add_window.geometry("400x400")

    tk.Label(add_window, text="Employee Name").pack(pady=5)
    name_entry = tk.Entry(add_window)
    name_entry.pack()

    tk.Label(add_window, text="Employee ID").pack(pady=5)
    id_entry = tk.Entry(add_window)
    id_entry.pack()

    tk.Label(add_window, text="Department").pack(pady=5)
    department_entry = tk.Entry(add_window)
    department_entry.pack()

    tk.Label(add_window, text="Position").pack(pady=5)
    position_entry = tk.Entry(add_window)
    position_entry.pack()

    tk.Label(add_window, text="Salary").pack(pady=5)
    salary_entry = tk.Entry(add_window)
    salary_entry.pack()

    def save_record():
        name = name_entry.get()
        employee_id = id_entry.get()
        department = department_entry.get()
        position = position_entry.get()
        salary = salary_entry.get()

        file_name = "employee_records.xlsx"

        if not os.path.exists(file_name):
            workbook = openpyxl.Workbook()
            sheet = workbook.active

            sheet.append([
                "Employee Name",
                "Employee ID",
                "Department",
                "Position",
                "Salary"
            ])

            workbook.save(file_name)

        workbook = openpyxl.load_workbook(file_name)
        sheet = workbook.active

        sheet.append([
            name,
            employee_id,
            department,
            position,
            salary
        ])

        workbook.save(file_name)

        messagebox.showinfo("Add Record", "Record saved successfully")
        add_window.destroy()

    tk.Button(
        add_window,
        text="Save",
        command=save_record
    ).pack(pady=20)


# ---------------- VIEW RECORDS ----------------

def view_records():
    file_name = "employee_records.xlsx"

    if not os.path.exists(file_name):
        messagebox.showinfo("View Records", "No records found")
        return

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    view_window = tk.Toplevel()
    view_window.title("View Employee Records")
    view_window.geometry("700x400")

    for row in sheet.iter_rows(values_only=True):
        row_text = " | ".join(str(value) for value in row)

        tk.Label(
            view_window,
            text=row_text
        ).pack(pady=3)


# ---------------- SEARCH ----------------

def search_record():
    file_name = "employee_records.xlsx"

    if not os.path.exists(file_name):
        messagebox.showinfo("Search", "No records found")
        return

    search_window = tk.Toplevel()
    search_window.title("Search Employee")
    search_window.geometry("400x300")

    tk.Label(
        search_window,
        text="Enter Employee ID"
    ).pack(pady=10)

    search_entry = tk.Entry(search_window)
    search_entry.pack()

    def search():
        employee_id = search_entry.get()

        workbook = openpyxl.load_workbook(file_name)
        sheet = workbook.active

        found = False

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if str(row[1]) == employee_id:

                result = (
                    "Employee Name: " + str(row[0]) + "\n"
                    "Employee ID: " + str(row[1]) + "\n"
                    "Department: " + str(row[2]) + "\n"
                    "Position: " + str(row[3]) + "\n"
                    "Salary: " + str(row[4])
                )

                tk.Label(
                    search_window,
                    text=result,
                    justify="left"
                ).pack(pady=20)

                found = True
                break

        if not found:
            messagebox.showinfo(
                "Search",
                "Employee not found"
            )

    tk.Button(
        search_window,
        text="Search",
        command=search
    ).pack(pady=15)


# ---------------- UPDATE RECORD ----------------

def update_record():
    file_name = "employee_records.xlsx"

    if not os.path.exists(file_name):
        messagebox.showinfo("Update", "No records found")
        return

    update_window = tk.Toplevel()
    update_window.title("Update Employee Record")
    update_window.geometry("400x400")

    tk.Label(update_window, text="Employee ID").pack(pady=5)
    id_entry = tk.Entry(update_window)
    id_entry.pack()

    tk.Label(update_window, text="Employee Name").pack(pady=5)
    name_entry = tk.Entry(update_window)
    name_entry.pack()

    tk.Label(update_window, text="Department").pack(pady=5)
    department_entry = tk.Entry(update_window)
    department_entry.pack()

    tk.Label(update_window, text="Position").pack(pady=5)
    position_entry = tk.Entry(update_window)
    position_entry.pack()

    tk.Label(update_window, text="Salary").pack(pady=5)
    salary_entry = tk.Entry(update_window)
    salary_entry.pack()

    def update():
        employee_id = id_entry.get()

        workbook = openpyxl.load_workbook(file_name)
        sheet = workbook.active

        found = False

        for row in sheet.iter_rows(min_row=2):

            if str(row[1].value) == employee_id:

                row[0].value = name_entry.get()
                row[2].value = department_entry.get()
                row[3].value = position_entry.get()
                row[4].value = salary_entry.get()

                found = True
                break

        if found:
            workbook.save(file_name)

            messagebox.showinfo(
                "Update",
                "Record updated successfully"
            )

            update_window.destroy()

        else:
            messagebox.showinfo(
                "Update",
                "Employee not found"
            )

    tk.Button(
        update_window,
        text="Update",
        command=update
    ).pack(pady=20)


# ---------------- DELETE RECORD ----------------

def delete_record():
    file_name = "employee_records.xlsx"

    if not os.path.exists(file_name):
        messagebox.showinfo("Delete", "No records found")
        return

    delete_window = tk.Toplevel()
    delete_window.title("Delete Employee Record")
    delete_window.geometry("400x250")

    tk.Label(
        delete_window,
        text="Enter Employee ID"
    ).pack(pady=15)

    id_entry = tk.Entry(delete_window)
    id_entry.pack()

    def delete():
        employee_id = id_entry.get()

        workbook = openpyxl.load_workbook(file_name)
        sheet = workbook.active

        found = False

        for row in range(2, sheet.max_row + 1):

            if str(sheet.cell(row, 2).value) == employee_id:

                sheet.delete_rows(row, 1)
                found = True
                break

        if found:
            workbook.save(file_name)

            messagebox.showinfo(
                "Delete",
                "Record deleted successfully"
            )

            delete_window.destroy()

        else:
            messagebox.showinfo(
                "Delete",
                "Employee not found"
            )

    tk.Button(
        delete_window,
        text="Delete",
        command=delete
    ).pack(pady=20)


# ---------------- DASHBOARD ----------------

def open_dashboard():
    root = tk.Tk()
    root.title("Employee Management System")
    root.geometry("500x500")

    title = tk.Label(
        root,
        text="Employee Management System",
        font=("Arial", 18)
    )
    title.pack(pady=30)

    tk.Button(
        root,
        text="Add Record",
        width=20,
        command=add_record
    ).pack(pady=5)

    tk.Button(
        root,
        text="View Records",
        width=20,
        command=view_records
    ).pack(pady=5)

    tk.Button(
        root,
        text="Search",
        width=20,
        command=search_record
    ).pack(pady=5)

    tk.Button(
        root,
        text="Update Record",
        width=20,
        command=update_record
    ).pack(pady=5)

    tk.Button(
        root,
        text="Delete Record",
        width=20,
        command=delete_record
    ).pack(pady=5)

    tk.Button(
        root,
        text="Logout",
        width=20,
        command=root.destroy
    ).pack(pady=20)

    root.mainloop()


# ---------------- LOGIN WINDOW ----------------

login_window = tk.Tk()
login_window.title("Employee Management System")
login_window.geometry("400x300")

tk.Label(
    login_window,
    text="Employee Management System",
    font=("Arial", 16)
).pack(pady=30)

tk.Label(
    login_window,
    text="Username"
).pack()

username_entry = tk.Entry(login_window)
username_entry.pack()

tk.Label(
    login_window,
    text="Password"
).pack()

password_entry = tk.Entry(
    login_window,
    show="*"
)
password_entry.pack()

tk.Button(
    login_window,
    text="Login",
    command=login
).pack(pady=20)

login_window.mainloop()
